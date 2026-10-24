from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

DARK_BLUE    = "1F3864"
MID_BLUE     = "2E75B6"
LIGHT_BLUE   = "D6E4F0"
GREEN_HDR    = "1E6B3C"
GREEN_LIGHT  = "D9EAD3"
RED_LIGHT    = "FCE4D6"
WHITE        = "FFFFFF"
GREY_ROW     = "F2F2F2"

def hdr_fill(h): return PatternFill("solid", fgColor=h)
def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)
def wrap(h="left", v="top"): return Alignment(wrap_text=True, horizontal=h, vertical=v)

def sh(cell, bg, fg=WHITE, size=11):
    cell.fill = hdr_fill(bg)
    cell.font = Font(name="Calibri", bold=True, color=fg, size=size)
    cell.alignment = wrap("center", "center")
    cell.border = thin_border()

def sb(cell, bg=WHITE, bold=False, align="left"):
    cell.fill = hdr_fill(bg)
    cell.font = Font(name="Calibri", bold=bold, size=10)
    cell.alignment = wrap(align, "top")
    cell.border = thin_border()

def title_row(ws, text, cols, row=1, bg=DARK_BLUE):
    ws.merge_cells(f"A{row}:{get_column_letter(cols)}{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.fill = hdr_fill(bg)
    c.font = Font(name="Calibri", bold=True, color=WHITE, size=14)
    c.alignment = wrap("center", "center")
    ws.row_dimensions[row].height = 30

def sub_row(ws, text, cols, row=2):
    ws.merge_cells(f"A{row}:{get_column_letter(cols)}{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.fill = hdr_fill(MID_BLUE)
    c.font = Font(name="Calibri", bold=True, color=WHITE, size=11)
    c.alignment = wrap("center", "center")
    ws.row_dimensions[row].height = 22

# ── SHEET 1: Test Scenarios ──────────────────────────────────────────────
ws1 = wb.active
ws1.title = "Test Scenarios"
title_row(ws1, "BULK BLOG SCHEDULING – TEST SCENARIOS", 6)
sub_row(ws1, "Feature: Schedule multiple approved blogs for one doctor with individual dates and times", 6)

hdrs = ["Scenario ID", "Test Scenario", "Description", "AC Covered", "Priority", "Linked Test Cases"]
for col, h in enumerate(hdrs, 1):
    sh(ws1.cell(row=3, column=col, value=h), MID_BLUE)
ws1.row_dimensions[3].height = 22

scenarios = [
    ("TS-01", "Multi-blog selection for one doctor",
     "Verify that a publishing user can select multiple approved blogs belonging to a single doctor from the blog repository.",
     "AC1", "High", "TC-001, TC-002, TC-003"),
    ("TS-02", "Single scheduling workspace for selected blogs",
     "Verify that all selected blogs open together in one unified scheduling workspace without requiring separate navigation per blog.",
     "AC2", "High", "TC-004, TC-005"),
    ("TS-03", "Individual date and time per blog",
     "Verify that each blog in the workspace has its own independently editable publication date and time field.",
     "AC3", "High", "TC-006, TC-007, TC-008"),
    ("TS-04", "Scheduling across different days weeks and months",
     "Verify that blogs can be assigned dates spanning different days, weeks, and months within the same scheduling session.",
     "AC4", "High", "TC-009, TC-010"),
    ("TS-05", "Validation before confirmation",
     "Verify that GC validates every selected blog has a valid future date and time before allowing the user to confirm the schedule.",
     "AC5", "Critical", "TC-011, TC-012, TC-013, TC-014, TC-015"),
    ("TS-06", "Individual scheduled status after confirmation",
     "Verify that after confirmation each blog reflects its own scheduled status with the correct assigned date and time.",
     "AC6", "High", "TC-016, TC-017, TC-018"),
    ("TS-07", "Cross-doctor blog selection restriction",
     "Verify that blogs from different doctors cannot be mixed in the same bulk scheduling session.",
     "AC1", "High", "TC-019, TC-020"),
    ("TS-08", "Non-approved blog exclusion",
     "Verify that only approved blogs are selectable for bulk scheduling; draft or rejected blogs are excluded.",
     "AC1", "High", "TC-021, TC-022"),
    ("TS-09", "Partial date entry and save behaviour",
     "Verify system behaviour when a user fills dates for some blogs but not all before attempting to confirm.",
     "AC5", "High", "TC-023, TC-024"),
    ("TS-10", "Large volume bulk scheduling",
     "Verify that the workspace handles scheduling 10 to 15 blogs simultaneously without performance or UI degradation.",
     "AC2, AC3", "Medium", "TC-025, TC-026"),
]

for i, row in enumerate(scenarios):
    r = i + 4
    bg = GREY_ROW if i % 2 == 0 else WHITE
    for col, val in enumerate(row, 1):
        sb(ws1.cell(row=r, column=col, value=val), bg, bold=(col == 1))
    ws1.row_dimensions[r].height = 55

for col, w in zip(range(1, 7), [14, 36, 58, 14, 12, 28]):
    ws1.column_dimensions[get_column_letter(col)].width = w
ws1.freeze_panes = "A4"

# ── SHEET 2: Test Cases ──────────────────────────────────────────────────
ws2 = wb.create_sheet("Test Cases")
title_row(ws2, "BULK BLOG SCHEDULING – DETAILED TEST CASES", 10)
sub_row(ws2, "Feature: Schedule multiple approved blogs for one doctor with individual dates and times", 10)

tc_hdrs = ["Scenario ID", "Test Case ID", "Test Case Title", "Type",
           "Preconditions", "Test Steps", "Expected Result", "AC", "Priority", "Status"]
for col, h in enumerate(tc_hdrs, 1):
    sh(ws2.cell(row=3, column=col, value=h), GREEN_HDR)
ws2.row_dimensions[3].height = 22

test_cases = [
    # TS-01
    ("TS-01","TC-001","Select multiple approved blogs for one doctor","Positive",
     "1. User logged in with publishing role\n2. Doctor has 3+ approved blogs in repository",
     "1. Navigate to Blog Repository\n2. Filter blogs by a specific doctor\n3. Select 3 approved blogs using checkboxes\n4. Observe selection state",
     "All 3 blogs are selected; a bulk action toolbar or button appears indicating multi-select is active",
     "AC1","High","Ready"),

    ("TS-01","TC-002","Select maximum blogs (15) for one doctor","Positive",
     "1. Doctor has 15+ approved blogs\n2. User has publishing role",
     "1. Filter by doctor\n2. Select 15 approved blogs\n3. Proceed to scheduling workspace",
     "All 15 blogs are selected and loaded into the scheduling workspace without error",
     "AC1","High","Ready"),

    ("TS-01","TC-003","Cannot select blogs from two different doctors simultaneously","Negative",
     "1. Approved blogs exist for Doctor A and Doctor B",
     "1. Select 2 approved blogs for Doctor A\n2. Attempt to also select a blog for Doctor B\n3. Observe system response",
     "System prevents cross-doctor selection; an informational message is shown (e.g. 'Bulk scheduling is limited to one doctor per session'); Doctor B blog is not added to selection",
     "AC1","High","Ready"),

    # TS-02
    ("TS-02","TC-004","Selected blogs open in one scheduling workspace","Positive",
     "1. User has selected 3+ approved blogs for one doctor",
     "1. Select 3 approved blogs for one doctor\n2. Click Schedule or Bulk Schedule action\n3. Observe the workspace that opens",
     "A single scheduling workspace opens listing all 3 selected blogs; user does not need to navigate separately for each blog",
     "AC2","High","Ready"),

    ("TS-02","TC-005","Workspace displays all selected blog titles and details","Positive",
     "1. 5 approved blogs selected for one doctor",
     "1. Open bulk scheduling workspace\n2. Verify each blog entry is visible\n3. Check that blog title and doctor name are shown per row",
     "All 5 blogs are listed in the workspace with their titles and the doctor name visible; no blog is missing from the list",
     "AC2","High","Ready"),

    # TS-03
    ("TS-03","TC-006","Each blog has its own editable date field","Positive",
     "1. Bulk scheduling workspace is open with 3 blogs",
     "1. Open bulk scheduling workspace\n2. Locate the date field for Blog 1\n3. Enter a date for Blog 1\n4. Locate the date field for Blog 2\n5. Enter a different date for Blog 2",
     "Each blog has an independent date input; changing Blog 1 date does not affect Blog 2 date",
     "AC3","High","Ready"),

    ("TS-03","TC-007","Each blog has its own editable time field","Positive",
     "1. Bulk scheduling workspace is open with 3 blogs",
     "1. Open bulk scheduling workspace\n2. Enter time 09:00 for Blog 1\n3. Enter time 14:30 for Blog 2\n4. Enter time 18:00 for Blog 3\n5. Verify each time is saved independently",
     "Each blog retains its own time value; no time value is shared or overwritten across blogs",
     "AC3","High","Ready"),

    ("TS-03","TC-008","Date and time fields are pre-populated or empty by default","Positive",
     "1. Bulk scheduling workspace is open",
     "1. Open workspace with 3 blogs\n2. Observe the initial state of date and time fields for each blog",
     "Date and time fields are either empty or show a placeholder; no blog has a pre-filled date that could cause accidental scheduling",
     "AC3","Medium","Ready"),

    # TS-04
    ("TS-04","TC-009","Blogs can be scheduled across different weeks","Positive",
     "1. Bulk scheduling workspace open with 4 blogs",
     "1. Assign Blog 1 a date in Week 1 of next month\n2. Assign Blog 2 a date in Week 2 of next month\n3. Assign Blog 3 a date in Week 3 of next month\n4. Assign Blog 4 a date in Week 4 of next month\n5. Confirm schedule",
     "All 4 blogs are scheduled on their respective dates across different weeks; no conflict or override occurs",
     "AC4","High","Ready"),

    ("TS-04","TC-010","Blogs can be scheduled across different months","Positive",
     "1. Bulk scheduling workspace open with 3 blogs",
     "1. Assign Blog 1 a date in the current month\n2. Assign Blog 2 a date 2 months ahead\n3. Assign Blog 3 a date 3 months ahead\n4. Confirm schedule",
     "All 3 blogs are scheduled on their respective future month dates; the publishing calendar reflects entries across all 3 months",
     "AC4","High","Ready"),

    # TS-05
    ("TS-05","TC-011","Validation blocks confirmation when one blog has no date","Negative",
     "1. Bulk scheduling workspace open with 3 blogs",
     "1. Enter valid date and time for Blog 1 and Blog 2\n2. Leave Blog 3 date and time empty\n3. Click Confirm",
     "System blocks confirmation; validation error highlights Blog 3 row (e.g. 'Date and time are required for all blogs'); no blog is scheduled",
     "AC5","Critical","Ready"),

    ("TS-05","TC-012","Validation blocks confirmation when one blog has no time","Negative",
     "1. Bulk scheduling workspace open with 3 blogs",
     "1. Enter valid date and time for Blog 1 and Blog 2\n2. Enter a date for Blog 3 but leave time empty\n3. Click Confirm",
     "System blocks confirmation; validation error shown for Blog 3 time field; no blog is scheduled until all times are filled",
     "AC5","Critical","Ready"),

    ("TS-05","TC-013","Validation blocks past date entry","Negative",
     "1. Bulk scheduling workspace open with 2 blogs",
     "1. Enter a past date (e.g. yesterday) for Blog 1\n2. Enter a valid future date for Blog 2\n3. Click Confirm",
     "System blocks confirmation; validation error on Blog 1 (e.g. 'Publication date must be a future date'); Blog 2 is not affected but schedule is not confirmed until all dates are valid",
     "AC5","Critical","Ready"),

    ("TS-05","TC-014","Validation blocks invalid date format","Negative",
     "1. Bulk scheduling workspace open",
     "1. Type an invalid date string (e.g. '32-13-2025' or 'abcd') in Blog 1 date field\n2. Click Confirm",
     "System shows a format validation error on Blog 1 date field; confirmation is blocked",
     "AC5","High","Ready"),

    ("TS-05","TC-015","Confirmation proceeds when all blogs have valid future dates and times","Positive",
     "1. Bulk scheduling workspace open with 3 blogs",
     "1. Enter valid future date and time for all 3 blogs\n2. Click Confirm",
     "Validation passes for all blogs; confirmation dialog or success message is shown; schedule is saved",
     "AC5","High","Ready"),

    # TS-06
    ("TS-06","TC-016","Each blog shows Scheduled status after confirmation","Positive",
     "1. 3 blogs successfully scheduled via bulk workspace",
     "1. Complete bulk scheduling for 3 blogs\n2. Navigate to Blog Repository or doctor's blog list\n3. Check status of each scheduled blog",
     "All 3 blogs display Scheduled status individually; each shows its own assigned date and time",
     "AC6","High","Ready"),

    ("TS-06","TC-017","Scheduled blogs appear in publishing calendar with correct dates","Positive",
     "1. 3 blogs scheduled across 3 different dates",
     "1. Navigate to Publishing Calendar\n2. Check the dates assigned to each blog",
     "Each blog appears on its correct scheduled date in the calendar; no blog appears on the wrong date",
     "AC6","High","Ready"),

    ("TS-06","TC-018","Scheduled status is independent per blog","Positive",
     "1. 3 blogs scheduled; Blog 2 is later rescheduled",
     "1. Reschedule Blog 2 to a new date\n2. Check status of Blog 1 and Blog 3",
     "Only Blog 2 status/date changes; Blog 1 and Blog 3 retain their original scheduled dates and statuses unchanged",
     "AC6","Medium","Ready"),

    # TS-07
    ("TS-07","TC-019","Blogs from different doctors cannot be bulk scheduled together","Negative",
     "1. Approved blogs exist for Doctor A and Doctor B",
     "1. Select Blog 1 (Doctor A) and Blog 2 (Doctor B)\n2. Attempt to open bulk scheduling workspace",
     "System prevents opening the workspace with mixed-doctor blogs; error message shown (e.g. 'Please select blogs for one doctor only')",
     "AC1","High","Ready"),

    ("TS-07","TC-020","Selecting a second doctor's blog deselects or warns user","Negative",
     "1. User has selected 2 blogs for Doctor A",
     "1. With Doctor A blogs selected attempt to check a blog belonging to Doctor B\n2. Observe system response",
     "System either prevents the selection with a tooltip/warning or automatically deselects Doctor A blogs and prompts the user to confirm the switch",
     "AC1","Medium","Ready"),

    # TS-08
    ("TS-08","TC-021","Draft blogs are not selectable for bulk scheduling","Negative",
     "1. Doctor has both approved and draft blogs",
     "1. Navigate to Blog Repository filtered by doctor\n2. Attempt to select a draft blog alongside approved blogs",
     "Draft blog checkbox is disabled or not shown; only approved blogs can be selected for bulk scheduling",
     "AC1","High","Ready"),

    ("TS-08","TC-022","Rejected blogs are not selectable for bulk scheduling","Negative",
     "1. Doctor has both approved and rejected blogs",
     "1. Navigate to Blog Repository filtered by doctor\n2. Attempt to select a rejected blog",
     "Rejected blog is not selectable; it is either hidden from the bulk selection view or its checkbox is disabled",
     "AC1","High","Ready"),

    # TS-09
    ("TS-09","TC-023","Confirmation blocked when only some blogs have dates filled","Negative",
     "1. Bulk scheduling workspace open with 5 blogs",
     "1. Fill valid date and time for Blog 1, 2, and 3\n2. Leave Blog 4 and Blog 5 empty\n3. Click Confirm",
     "System blocks confirmation and highlights Blog 4 and Blog 5 rows with validation errors; user must fill all entries before confirming",
     "AC5","High","Ready"),

    ("TS-09","TC-024","User can save progress and return to complete scheduling","Positive",
     "1. Bulk scheduling workspace open with 5 blogs",
     "1. Fill dates for Blog 1 and Blog 2\n2. Click Save Draft or equivalent\n3. Close workspace\n4. Reopen the scheduling session\n5. Verify previously entered dates are retained",
     "Previously entered dates for Blog 1 and Blog 2 are retained; user can continue filling remaining blogs without re-entering data",
     "AC3","Medium","Ready"),

    # TS-10
    ("TS-10","TC-025","Workspace loads 15 blogs without performance issues","Positive",
     "1. Doctor has 15 approved blogs\n2. All 15 are selected",
     "1. Select all 15 approved blogs for one doctor\n2. Open bulk scheduling workspace\n3. Measure load time and UI responsiveness",
     "Workspace loads within acceptable time (e.g. under 5 seconds); all 15 blog rows are visible and scrollable; no UI freeze or error",
     "AC2, AC3","Medium","Ready"),

    ("TS-10","TC-026","All 15 blogs can be individually dated and confirmed","Positive",
     "1. Bulk scheduling workspace open with 15 blogs",
     "1. Enter a unique future date and time for each of the 15 blogs\n2. Click Confirm",
     "All 15 blogs pass validation and are confirmed with their individual scheduled dates; each appears with Scheduled status in the repository",
     "AC3, AC5, AC6","Medium","Ready"),
]

for i, tc in enumerate(test_cases):
    row = i + 4
    bg = GREEN_LIGHT if tc[3] == "Positive" else RED_LIGHT
    for col, val in enumerate(tc, 1):
        c = ws2.cell(row=row, column=col, value=val)
        sb(c, bg, bold=(col in (1, 2)), align="center" if col == 4 else "left")
    ws2.row_dimensions[row].height = 90

for col, w in zip(range(1, 11), [12, 12, 40, 12, 40, 58, 48, 10, 10, 10]):
    ws2.column_dimensions[get_column_letter(col)].width = w
ws2.freeze_panes = "A4"

# ── SHEET 3: Traceability Matrix ─────────────────────────────────────────
ws3 = wb.create_sheet("Scenario-TC Mapping")
title_row(ws3, "SCENARIO → TEST CASE TRACEABILITY MATRIX", 7)
sub_row(ws3, "Feature: Schedule multiple approved blogs for one doctor with individual dates and times", 7)

map_hdrs = ["Scenario ID", "Test Scenario", "Test Case ID", "Test Case Title", "Type", "AC", "Priority"]
for col, h in enumerate(map_hdrs, 1):
    sh(ws3.cell(row=3, column=col, value=h), DARK_BLUE)
ws3.row_dimensions[3].height = 22

mapping = [
    ("TS-01","Multi-blog selection for one doctor","TC-001","Select multiple approved blogs for one doctor","Positive","AC1","High"),
    ("TS-01","Multi-blog selection for one doctor","TC-002","Select maximum blogs (15) for one doctor","Positive","AC1","High"),
    ("TS-01","Multi-blog selection for one doctor","TC-003","Cannot select blogs from two different doctors simultaneously","Negative","AC1","High"),
    ("TS-02","Single scheduling workspace for selected blogs","TC-004","Selected blogs open in one scheduling workspace","Positive","AC2","High"),
    ("TS-02","Single scheduling workspace for selected blogs","TC-005","Workspace displays all selected blog titles and details","Positive","AC2","High"),
    ("TS-03","Individual date and time per blog","TC-006","Each blog has its own editable date field","Positive","AC3","High"),
    ("TS-03","Individual date and time per blog","TC-007","Each blog has its own editable time field","Positive","AC3","High"),
    ("TS-03","Individual date and time per blog","TC-008","Date and time fields are pre-populated or empty by default","Positive","AC3","Medium"),
    ("TS-04","Scheduling across different days weeks and months","TC-009","Blogs can be scheduled across different weeks","Positive","AC4","High"),
    ("TS-04","Scheduling across different days weeks and months","TC-010","Blogs can be scheduled across different months","Positive","AC4","High"),
    ("TS-05","Validation before confirmation","TC-011","Validation blocks confirmation when one blog has no date","Negative","AC5","Critical"),
    ("TS-05","Validation before confirmation","TC-012","Validation blocks confirmation when one blog has no time","Negative","AC5","Critical"),
    ("TS-05","Validation before confirmation","TC-013","Validation blocks past date entry","Negative","AC5","Critical"),
    ("TS-05","Validation before confirmation","TC-014","Validation blocks invalid date format","Negative","AC5","High"),
    ("TS-05","Validation before confirmation","TC-015","Confirmation proceeds when all blogs have valid future dates and times","Positive","AC5","High"),
    ("TS-06","Individual scheduled status after confirmation","TC-016","Each blog shows Scheduled status after confirmation","Positive","AC6","High"),
    ("TS-06","Individual scheduled status after confirmation","TC-017","Scheduled blogs appear in publishing calendar with correct dates","Positive","AC6","High"),
    ("TS-06","Individual scheduled status after confirmation","TC-018","Scheduled status is independent per blog","Positive","AC6","Medium"),
    ("TS-07","Cross-doctor blog selection restriction","TC-019","Blogs from different doctors cannot be bulk scheduled together","Negative","AC1","High"),
    ("TS-07","Cross-doctor blog selection restriction","TC-020","Selecting a second doctor's blog deselects or warns user","Negative","AC1","Medium"),
    ("TS-08","Non-approved blog exclusion","TC-021","Draft blogs are not selectable for bulk scheduling","Negative","AC1","High"),
    ("TS-08","Non-approved blog exclusion","TC-022","Rejected blogs are not selectable for bulk scheduling","Negative","AC1","High"),
    ("TS-09","Partial date entry and save behaviour","TC-023","Confirmation blocked when only some blogs have dates filled","Negative","AC5","High"),
    ("TS-09","Partial date entry and save behaviour","TC-024","User can save progress and return to complete scheduling","Positive","AC3","Medium"),
    ("TS-10","Large volume bulk scheduling","TC-025","Workspace loads 15 blogs without performance issues","Positive","AC2, AC3","Medium"),
    ("TS-10","Large volume bulk scheduling","TC-026","All 15 blogs can be individually dated and confirmed","Positive","AC3, AC5, AC6","Medium"),
]

prev_sid = None
toggle = True
for i, row_data in enumerate(mapping):
    row = i + 4
    if row_data[0] != prev_sid:
        toggle = not toggle
        prev_sid = row_data[0]
    bg = LIGHT_BLUE if toggle else WHITE
    type_bg = GREEN_LIGHT if row_data[4] == "Positive" else RED_LIGHT
    for col, val in enumerate(row_data, 1):
        c = ws3.cell(row=row, column=col, value=val)
        if col == 5:
            sb(c, type_bg, align="center")
        elif col in (1, 3):
            sb(c, bg, bold=True, align="center")
        else:
            sb(c, bg)
    ws3.row_dimensions[row].height = 40

for col, w in zip(range(1, 8), [14, 40, 12, 50, 12, 16, 12]):
    ws3.column_dimensions[get_column_letter(col)].width = w
ws3.freeze_panes = "A4"

# ── SHEET 4: AC Coverage ─────────────────────────────────────────────────
ws4 = wb.create_sheet("AC Coverage")
title_row(ws4, "ACCEPTANCE CRITERIA COVERAGE MATRIX", 5)
ws4.row_dimensions[1].height = 30

ac_hdrs = ["AC ID", "Acceptance Criteria Statement", "Covered By (Test Cases)", "Total TCs", "Coverage"]
for col, h in enumerate(ac_hdrs, 1):
    sh(ws4.cell(row=2, column=col, value=h), MID_BLUE)
ws4.row_dimensions[2].height = 22

ac_data = [
    ("AC1","The user can select multiple approved blogs belonging to one doctor",
     "TC-001, TC-002, TC-003, TC-019, TC-020, TC-021, TC-022", 7, "Covered"),
    ("AC2","The selected blogs open in one scheduling workspace",
     "TC-004, TC-005, TC-025", 3, "Covered"),
    ("AC3","Each blog has its own editable publication date and time",
     "TC-006, TC-007, TC-008, TC-024, TC-025, TC-026", 6, "Covered"),
    ("AC4","The user can schedule blogs across different days, weeks or months",
     "TC-009, TC-010", 2, "Covered"),
    ("AC5","GC validates that every selected blog has a valid date and time before confirmation",
     "TC-011, TC-012, TC-013, TC-014, TC-015, TC-023, TC-026", 7, "Covered"),
    ("AC6","After confirmation, each blog appears with its own scheduled status",
     "TC-016, TC-017, TC-018, TC-026", 4, "Covered"),
]

for i, row_data in enumerate(ac_data):
    row = i + 3
    bg = GREY_ROW if i % 2 == 0 else WHITE
    for col, val in enumerate(row_data, 1):
        c = ws4.cell(row=row, column=col, value=val)
        align = "center" if col in (1, 4, 5) else "left"
        sb(c, bg, bold=(col == 1), align=align)
        if col == 5:
            c.font = Font(name="Calibri", bold=True, color="1E6B3C", size=10)
    ws4.row_dimensions[row].height = 55

for col, w in zip(range(1, 6), [10, 65, 55, 12, 12]):
    ws4.column_dimensions[get_column_letter(col)].width = w
ws4.freeze_panes = "A3"

out = r"c:\Users\DeepthiKonchada\Downloads\AutomationNew\testData\BulkScheduling_TestCases.xlsx"
wb.save(out)
print(f"Created: {out}")
