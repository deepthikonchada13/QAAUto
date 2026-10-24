from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

wb = Workbook()

# ── colour palette ──────────────────────────────────────────────────────────
DARK_BLUE   = "1F3864"
MID_BLUE    = "2E75B6"
LIGHT_BLUE  = "D6E4F0"
GREEN_HDR   = "1E6B3C"
GREEN_LIGHT = "D9EAD3"
RED_LIGHT   = "FCE4D6"
YELLOW_LIGHT= "FFF2CC"
WHITE       = "FFFFFF"
GREY_ROW    = "F2F2F2"

def hdr_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def cell_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def hdr_font(color=WHITE, size=11, bold=True):
    return Font(name="Calibri", bold=bold, color=color, size=size)

def body_font(bold=False, size=10):
    return Font(name="Calibri", bold=bold, size=size)

def wrap_align(horizontal="left", vertical="top"):
    return Alignment(wrap_text=True, horizontal=horizontal, vertical=vertical)

def style_header_cell(cell, bg, fg=WHITE, size=11):
    cell.fill = hdr_fill(bg)
    cell.font = hdr_font(fg, size)
    cell.alignment = wrap_align("center", "center")
    cell.border = thin_border()

def style_body_cell(cell, bg=WHITE, bold=False, align="left"):
    cell.fill = cell_fill(bg)
    cell.font = body_font(bold)
    cell.alignment = wrap_align(align, "top")
    cell.border = thin_border()

# ════════════════════════════════════════════════════════════════════════════
# SHEET 1 – TEST SCENARIOS
# ════════════════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Test Scenarios"

# Title banner
ws1.merge_cells("A1:F1")
ws1["A1"] = "MARK AS USED – TEST SCENARIOS"
ws1["A1"].fill = hdr_fill(DARK_BLUE)
ws1["A1"].font = Font(name="Calibri", bold=True, color=WHITE, size=14)
ws1["A1"].alignment = wrap_align("center", "center")
ws1.row_dimensions[1].height = 30

# Sub-header
ws1.merge_cells("A2:F2")
ws1["A2"] = "Feature: Manual Mark as Used for blogs published outside GC"
ws1["A2"].fill = hdr_fill(MID_BLUE)
ws1["A2"].font = Font(name="Calibri", bold=True, color=WHITE, size=11)
ws1["A2"].alignment = wrap_align("center", "center")
ws1.row_dimensions[2].height = 22

# Column headers
ts_headers = ["Scenario ID", "Test Scenario", "Description", "Acceptance Criteria Covered", "Priority", "Linked Test Cases"]
for col, h in enumerate(ts_headers, 1):
    c = ws1.cell(row=3, column=col, value=h)
    style_header_cell(c, MID_BLUE)
ws1.row_dimensions[3].height = 22

scenarios = [
    ("TS-01", "Auto status update via GC tracking",
     "Verify that when a blog is published through a GC-connected website, GC automatically marks it as Used and removes it from the unused list.",
     "AC1, AC2, AC3", "High", "TC-001, TC-002, TC-003"),

    ("TS-02", "Manual Mark as Used for outside-GC publications",
     "Verify that a user can manually mark a blog as Used and map it to a customer website when the blog was published outside GC.",
     "AC3, AC6", "High", "TC-004, TC-005, TC-006"),

    ("TS-03", "Audit trail for manual Mark as Used",
     "Verify that every manual Mark as Used action is recorded in the audit trail with user, timestamp, blog name, and mapped website.",
     "AC5", "High", "TC-007, TC-008"),

    ("TS-04", "Prevent duplicate manual marking",
     "Verify that a blog with an existing GC publication record cannot be manually marked as Used again.",
     "AC4", "Critical", "TC-009, TC-010"),

    ("TS-05", "Mark as Used unavailability for ineligible blogs",
     "Verify that the Mark as Used option is not available for blogs that were never downloaded outside GC or are already in Used status.",
     "AC3, AC4", "Medium", "TC-011, TC-012"),

    ("TS-06", "Mandatory website mapping validation",
     "Verify that the Mark as Used action cannot be completed without selecting or entering a valid customer website.",
     "AC3", "High", "TC-013, TC-014"),

    ("TS-07", "Role-based access control for Mark as Used",
     "Verify that only users with the publishing role can perform the manual Mark as Used action.",
     "AC3", "High", "TC-015, TC-016"),

    ("TS-08", "GC auto-tracking scope boundary",
     "Verify that GC auto-tracking applies only to connected websites and that non-connected website publications require manual marking.",
     "AC1, AC2, AC3", "High", "TC-017, TC-018"),

    ("TS-09", "Fallback when GC tracking fails",
     "Verify that when GC fails to capture a publication status, the manual Mark as Used option is available as a fallback.",
     "AC3, AC5", "High", "TC-019, TC-020"),
]

for i, (sid, title, desc, ac, pri, linked) in enumerate(scenarios):
    row = i + 4
    bg = GREY_ROW if i % 2 == 0 else WHITE
    data = [sid, title, desc, ac, pri, linked]
    for col, val in enumerate(data, 1):
        c = ws1.cell(row=row, column=col, value=val)
        style_body_cell(c, bg, bold=(col == 1))
    ws1.row_dimensions[row].height = 55

# Column widths
for col, w in zip(range(1, 7), [14, 32, 55, 28, 12, 28]):
    ws1.column_dimensions[get_column_letter(col)].width = w

ws1.freeze_panes = "A4"

# ════════════════════════════════════════════════════════════════════════════
# SHEET 2 – TEST CASES
# ════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Test Cases")

ws2.merge_cells("A1:J1")
ws2["A1"] = "MARK AS USED – DETAILED TEST CASES"
ws2["A1"].fill = hdr_fill(DARK_BLUE)
ws2["A1"].font = Font(name="Calibri", bold=True, color=WHITE, size=14)
ws2["A1"].alignment = wrap_align("center", "center")
ws2.row_dimensions[1].height = 30

ws2.merge_cells("A2:J2")
ws2["A2"] = "Feature: Manual Mark as Used for blogs published outside GC"
ws2["A2"].fill = hdr_fill(MID_BLUE)
ws2["A2"].font = Font(name="Calibri", bold=True, color=WHITE, size=11)
ws2["A2"].alignment = wrap_align("center", "center")
ws2.row_dimensions[2].height = 22

tc_headers = [
    "Scenario ID", "Test Case ID", "Test Case Title", "Type",
    "Preconditions", "Test Steps", "Expected Result",
    "Acceptance Criteria", "Priority", "Status"
]
for col, h in enumerate(tc_headers, 1):
    c = ws2.cell(row=3, column=col, value=h)
    style_header_cell(c, GREEN_HDR)
ws2.row_dimensions[3].height = 22

test_cases = [
    # TS-01
    ("TS-01","TC-001","Auto mark blog as Used after successful GC-tracked publication","Positive",
     "1. User logged in with publishing role\n2. Blog exists in unused repository\n3. External website connected to GC",
     "1. Navigate to Blog Repository\n2. Select an unused blog\n3. Click Publish via GC\n4. Select a GC-connected website\n5. Complete the publication flow\n6. Observe blog status after successful publication",
     "Blog status is automatically updated to Used without any manual action required",
     "AC1","High","Ready"),

    ("TS-01","TC-002","Auto-remove blog from unused repository after GC-tracked publication","Positive",
     "1. User logged in with publishing role\n2. Blog exists in unused repository\n3. External website connected to GC",
     "1. Navigate to Blog Repository > Unused Blogs list\n2. Note the blog count\n3. Publish a blog via GC to a connected website\n4. Wait for publication confirmation\n5. Return to Unused Blogs list\n6. Verify the published blog is no longer listed",
     "Published blog is removed from the unused repository list and count decreases by 1",
     "AC2","High","Ready"),

    ("TS-01","TC-003","Verify Mark as Used option is NOT shown for GC-tracked publications","Positive",
     "1. Blog has been published and tracked through GC\n2. Blog status is Used",
     "1. Navigate to Blog Repository\n2. Locate the blog published via GC\n3. Open the blog action menu or detail view\n4. Check for presence of Mark as Used option",
     "Mark as Used option is not visible or is disabled for blogs already tracked by GC",
     "AC3","High","Ready"),

    # TS-02
    ("TS-02","TC-004","Manual Mark as Used available for blog downloaded and published outside GC","Positive",
     "1. User has publishing role\n2. Blog was downloaded from GC\n3. Blog was published on a non-GC-connected website\n4. Blog still shows as unused in GC",
     "1. Navigate to Blog Repository > Unused Blogs\n2. Locate the downloaded blog\n3. Open the blog action menu\n4. Verify Mark as Used option is present\n5. Click Mark as Used\n6. Select the customer website\n7. Confirm the action",
     "Mark as Used option is available; blog status updates to Used and is removed from unused list",
     "AC3","High","Ready"),

    ("TS-02","TC-005","User can map blog to customer website during manual Mark as Used","Positive",
     "1. Blog was published outside GC\n2. Blog is in unused state in GC",
     "1. Click Mark as Used on the unused blog\n2. Verify a website mapping field is present in the dialog\n3. Select or enter the customer website URL\n4. Submit the form",
     "Blog is marked as Used and mapped to the specified customer website; mapping is saved and visible in blog details",
     "AC3, AC6","High","Ready"),

    ("TS-02","TC-006","Repository and publishing counts update after manual Mark as Used","Positive",
     "1. Blog is in unused state\n2. User has publishing role",
     "1. Note current unused blog count and publishing count\n2. Manually mark a blog as Used\n3. Map it to a customer website\n4. Confirm the action\n5. Check unused blog count\n6. Check publishing count",
     "Unused blog count decreases by 1; publishing count increases by 1; changes reflected immediately",
     "AC6","High","Ready"),

    # TS-03
    ("TS-03","TC-007","Manual Mark as Used action is recorded in audit trail","Positive",
     "1. User logged in with publishing role\n2. Blog is in unused state",
     "1. Manually mark a blog as Used and map to a website\n2. Navigate to Audit Trail / Activity Log\n3. Search or filter by blog name or action type\n4. Locate the audit entry for the Mark as Used action",
     "Audit trail contains: user who performed the action, timestamp, blog name, action type (Mark as Used), and mapped website",
     "AC5","High","Ready"),

    ("TS-03","TC-008","Audit trail entry contains correct user and timestamp","Positive",
     "1. User A is logged in\n2. Blog is in unused state",
     "1. Log in as User A\n2. Manually mark a blog as Used\n3. Navigate to Audit Trail\n4. Locate the entry for the action",
     "Audit entry shows the correct username (User A) and an accurate timestamp matching the time of the action",
     "AC5","Medium","Ready"),

    # TS-04
    ("TS-04","TC-009","Cannot manually mark as Used if GC publication record already exists","Negative",
     "1. Blog has an existing successful GC publication record\n2. Blog status is Used (auto-tracked)",
     "1. Navigate to Blog Repository\n2. Locate a blog already published and tracked via GC\n3. Attempt to access Mark as Used option\n4. Try to submit the manual marking",
     "System prevents the action; error message displayed stating the blog already has a GC publication record; status remains unchanged",
     "AC4","Critical","Ready"),

    ("TS-04","TC-010","Error message shown when attempting duplicate manual Mark as Used","Negative",
     "1. Blog already has a GC-tracked Used status",
     "1. Attempt to manually mark the blog as Used via any available UI path\n2. Observe the system response",
     "Clear error message displayed (e.g. 'This blog has already been published and tracked through GC. Manual marking is not allowed.'); no duplicate record created",
     "AC4","High","Ready"),

    # TS-05
    ("TS-05","TC-011","Mark as Used not available for blogs never downloaded from GC","Negative",
     "1. Blog exists in repository\n2. Blog has never been downloaded or published",
     "1. Navigate to Blog Repository\n2. Select a blog that has never been downloaded\n3. Open the action menu\n4. Check for Mark as Used option",
     "Mark as Used option is not available for blogs with no download history outside GC",
     "AC3","Medium","Ready"),

    ("TS-05","TC-012","Mark as Used not available for blogs already in Used status","Negative",
     "1. Blog is already in Used status (any source)",
     "1. Navigate to Blog Repository\n2. Locate a blog with Used status\n3. Open the action menu\n4. Check for Mark as Used option",
     "Mark as Used option is not available or is disabled for blogs already in Used status",
     "AC3, AC4","Medium","Ready"),

    # TS-06
    ("TS-06","TC-013","Cannot submit Mark as Used without selecting a customer website","Negative",
     "1. Blog is in unused state\n2. Mark as Used dialog is open",
     "1. Open Mark as Used dialog for an unused blog\n2. Leave the customer website field empty\n3. Click Confirm or Submit",
     "System displays a validation error (e.g. 'Please select or enter the customer website'); form is not submitted; blog status remains unchanged",
     "AC3","High","Ready"),

    ("TS-06","TC-014","Cannot submit Mark as Used with invalid website URL","Negative",
     "1. Blog is in unused state\n2. Mark as Used dialog is open",
     "1. Open Mark as Used dialog\n2. Enter an invalid or malformed URL (e.g. 'notawebsite' or 'http://')\n3. Click Confirm or Submit",
     "System displays a validation error for invalid URL format; form is not submitted",
     "AC3","Medium","Ready"),

    # TS-07
    ("TS-07","TC-015","Publishing user can perform manual Mark as Used","Positive",
     "1. User has publishing role\n2. Blog is in unused state and eligible for manual marking",
     "1. Log in as a user with publishing role\n2. Navigate to unused blog\n3. Attempt to Mark as Used and map to website",
     "Action is allowed; blog is marked as Used successfully",
     "AC3","High","Ready"),

    ("TS-07","TC-016","Non-publishing user cannot perform manual Mark as Used","Negative",
     "1. User has read-only or non-publishing role\n2. Blog is in unused state",
     "1. Log in as a user without publishing role\n2. Navigate to unused blog\n3. Check for Mark as Used option",
     "Mark as Used option is not visible or is disabled for non-publishing users; action cannot be performed",
     "AC3","High","Ready"),

    # TS-08
    ("TS-08","TC-017","Blog published to non-GC-connected website remains unused until manually marked","Positive",
     "1. Blog is in unused state\n2. Target website is NOT connected to GC",
     "1. Download the blog from GC\n2. Manually publish it on a non-GC-connected website\n3. Return to GC and check blog status",
     "Blog status remains Unused in GC; no automatic update occurs; manual Mark as Used option is available",
     "AC3","High","Ready"),

    ("TS-08","TC-018","GC auto-tracks and marks Used only for connected website publications","Positive",
     "1. Blog is in unused state\n2. Target website IS connected to GC",
     "1. Publish blog via GC to a connected website\n2. Verify GC receives publication confirmation\n3. Check blog status in repository",
     "Blog is automatically marked as Used; removed from unused list; no manual action needed",
     "AC1, AC2","High","Ready"),

    # TS-09
    ("TS-09","TC-019","Mark as Used available when GC publication tracking fails","Positive",
     "1. Blog was published via GC but tracking failed (e.g. network error or webhook failure)\n2. Blog still shows as unused",
     "1. Identify a blog where GC publication tracking failed\n2. Navigate to the blog in repository\n3. Check for Mark as Used option",
     "Mark as Used option is available as a fallback; user can manually update the status and map to the website",
     "AC3","High","Ready"),

    ("TS-09","TC-020","Audit trail distinguishes between auto-tracked and manually marked Used status","Positive",
     "1. One blog was auto-tracked by GC\n2. Another blog was manually marked as Used",
     "1. Navigate to Audit Trail\n2. Compare entries for auto-tracked vs manually marked blogs",
     "Audit entries clearly indicate the source of Used status: System (auto-tracked) vs User (manual); timestamps and user details accurate for manual entries",
     "AC5","Medium","Ready"),
]

for i, tc in enumerate(test_cases):
    row = i + 4
    tc_type = tc[3]
    if tc_type == "Positive":
        row_bg = GREEN_LIGHT
    elif tc_type == "Negative":
        row_bg = RED_LIGHT
    else:
        row_bg = WHITE

    for col, val in enumerate(tc, 1):
        c = ws2.cell(row=row, column=col, value=val)
        is_id_col = col in (1, 2)
        is_type_col = col == 4
        style_body_cell(c, row_bg, bold=is_id_col)
        if is_type_col:
            c.alignment = wrap_align("center", "top")
    ws2.row_dimensions[row].height = 90

# Column widths
for col, w in zip(range(1, 11), [12, 12, 38, 12, 38, 55, 45, 14, 10, 10]):
    ws2.column_dimensions[get_column_letter(col)].width = w

ws2.freeze_panes = "A4"

# ════════════════════════════════════════════════════════════════════════════
# SHEET 3 – SCENARIO-TO-TESTCASE MAPPING
# ════════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Scenario-TC Mapping")

ws3.merge_cells("A1:G1")
ws3["A1"] = "TEST SCENARIO → TEST CASE TRACEABILITY MATRIX"
ws3["A1"].fill = hdr_fill(DARK_BLUE)
ws3["A1"].font = Font(name="Calibri", bold=True, color=WHITE, size=14)
ws3["A1"].alignment = wrap_align("center", "center")
ws3.row_dimensions[1].height = 30

ws3.merge_cells("A2:G2")
ws3["A2"] = "Feature: Manual Mark as Used for blogs published outside GC"
ws3["A2"].fill = hdr_fill(MID_BLUE)
ws3["A2"].font = Font(name="Calibri", bold=True, color=WHITE, size=11)
ws3["A2"].alignment = wrap_align("center", "center")
ws3.row_dimensions[2].height = 22

map_headers = ["Scenario ID", "Test Scenario", "Test Case ID", "Test Case Title", "Type", "Acceptance Criteria", "Priority"]
for col, h in enumerate(map_headers, 1):
    c = ws3.cell(row=3, column=col, value=h)
    style_header_cell(c, DARK_BLUE)
ws3.row_dimensions[3].height = 22

mapping_rows = [
    ("TS-01","Auto status update via GC tracking","TC-001","Auto mark blog as Used after successful GC-tracked publication","Positive","AC1","High"),
    ("TS-01","Auto status update via GC tracking","TC-002","Auto-remove blog from unused repository after GC-tracked publication","Positive","AC2","High"),
    ("TS-01","Auto status update via GC tracking","TC-003","Verify Mark as Used option is NOT shown for GC-tracked publications","Positive","AC3","High"),
    ("TS-02","Manual Mark as Used for outside-GC publications","TC-004","Manual Mark as Used available for blog downloaded and published outside GC","Positive","AC3","High"),
    ("TS-02","Manual Mark as Used for outside-GC publications","TC-005","User can map blog to customer website during manual Mark as Used","Positive","AC3, AC6","High"),
    ("TS-02","Manual Mark as Used for outside-GC publications","TC-006","Repository and publishing counts update after manual Mark as Used","Positive","AC6","High"),
    ("TS-03","Audit trail for manual Mark as Used","TC-007","Manual Mark as Used action is recorded in audit trail","Positive","AC5","High"),
    ("TS-03","Audit trail for manual Mark as Used","TC-008","Audit trail entry contains correct user and timestamp","Positive","AC5","Medium"),
    ("TS-04","Prevent duplicate manual marking","TC-009","Cannot manually mark as Used if GC publication record already exists","Negative","AC4","Critical"),
    ("TS-04","Prevent duplicate manual marking","TC-010","Error message shown when attempting duplicate manual Mark as Used","Negative","AC4","High"),
    ("TS-05","Mark as Used unavailability for ineligible blogs","TC-011","Mark as Used not available for blogs never downloaded from GC","Negative","AC3","Medium"),
    ("TS-05","Mark as Used unavailability for ineligible blogs","TC-012","Mark as Used not available for blogs already in Used status","Negative","AC3, AC4","Medium"),
    ("TS-06","Mandatory website mapping validation","TC-013","Cannot submit Mark as Used without selecting a customer website","Negative","AC3","High"),
    ("TS-06","Mandatory website mapping validation","TC-014","Cannot submit Mark as Used with invalid website URL","Negative","AC3","Medium"),
    ("TS-07","Role-based access control for Mark as Used","TC-015","Publishing user can perform manual Mark as Used","Positive","AC3","High"),
    ("TS-07","Role-based access control for Mark as Used","TC-016","Non-publishing user cannot perform manual Mark as Used","Negative","AC3","High"),
    ("TS-08","GC auto-tracking scope boundary","TC-017","Blog published to non-GC-connected website remains unused until manually marked","Positive","AC3","High"),
    ("TS-08","GC auto-tracking scope boundary","TC-018","GC auto-tracks and marks Used only for connected website publications","Positive","AC1, AC2","High"),
    ("TS-09","Fallback when GC tracking fails","TC-019","Mark as Used available when GC publication tracking fails","Positive","AC3","High"),
    ("TS-09","Fallback when GC tracking fails","TC-020","Audit trail distinguishes between auto-tracked and manually marked Used status","Positive","AC5","Medium"),
]

prev_scenario = None
scenario_start = 4
scenario_color_toggle = True

for i, row_data in enumerate(mapping_rows):
    row = i + 4
    scenario_id = row_data[0]

    if scenario_id != prev_scenario:
        scenario_color_toggle = not scenario_color_toggle
        prev_scenario = scenario_id

    bg = LIGHT_BLUE if scenario_color_toggle else WHITE
    tc_type = row_data[4]
    type_bg = GREEN_LIGHT if tc_type == "Positive" else RED_LIGHT

    for col, val in enumerate(row_data, 1):
        c = ws3.cell(row=row, column=col, value=val)
        if col == 5:
            style_body_cell(c, type_bg, bold=False, align="center")
        elif col in (1, 3):
            style_body_cell(c, bg, bold=True, align="center")
        else:
            style_body_cell(c, bg)
    ws3.row_dimensions[row].height = 40

for col, w in zip(range(1, 8), [14, 38, 12, 48, 12, 20, 12]):
    ws3.column_dimensions[get_column_letter(col)].width = w

ws3.freeze_panes = "A4"

# ════════════════════════════════════════════════════════════════════════════
# SHEET 4 – ACCEPTANCE CRITERIA COVERAGE
# ════════════════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("AC Coverage")

ws4.merge_cells("A1:E1")
ws4["A1"] = "ACCEPTANCE CRITERIA COVERAGE MATRIX"
ws4["A1"].fill = hdr_fill(DARK_BLUE)
ws4["A1"].font = Font(name="Calibri", bold=True, color=WHITE, size=14)
ws4["A1"].alignment = wrap_align("center", "center")
ws4.row_dimensions[1].height = 30

ac_headers = ["AC ID", "Acceptance Criteria Statement", "Covered By (Test Cases)", "Total TCs", "Coverage"]
for col, h in enumerate(ac_headers, 1):
    c = ws4.cell(row=2, column=col, value=h)
    style_header_cell(c, MID_BLUE)
ws4.row_dimensions[2].height = 22

ac_data = [
    ("AC1", "When a blog is successfully published and tracked through GC, its status is automatically updated to Used.",
     "TC-001, TC-018", 2, "✅ Covered"),
    ("AC2", "A blog successfully tracked through GC is automatically removed from the unused repository list.",
     "TC-002, TC-018", 2, "✅ Covered"),
    ("AC3", "The manual Mark as Used option is available only for publications completed outside GC or for publications that GC could not track.",
     "TC-003, TC-004, TC-005, TC-011, TC-012, TC-013, TC-014, TC-015, TC-016, TC-017, TC-019", 11, "✅ Covered"),
    ("AC4", "If a successful GC publication record already exists, the same blog cannot be manually marked as Used again.",
     "TC-009, TC-010, TC-012", 3, "✅ Covered"),
    ("AC5", "The manual action is recorded in the audit trail.",
     "TC-007, TC-008, TC-020", 3, "✅ Covered"),
    ("AC6", "The updated status is reflected in repository and publishing counts.",
     "TC-005, TC-006", 2, "✅ Covered"),
]

for i, row_data in enumerate(ac_data):
    row = i + 3
    bg = GREY_ROW if i % 2 == 0 else WHITE
    for col, val in enumerate(row_data, 1):
        c = ws4.cell(row=row, column=col, value=val)
        bold = col == 1
        align = "center" if col in (1, 4, 5) else "left"
        style_body_cell(c, bg, bold=bold, align=align)
        if col == 5:
            c.font = Font(name="Calibri", bold=True, color="1E6B3C", size=10)
    ws4.row_dimensions[row].height = 55

for col, w in zip(range(1, 6), [10, 65, 55, 12, 14]):
    ws4.column_dimensions[get_column_letter(col)].width = w

ws4.freeze_panes = "A3"

# ════════════════════════════════════════════════════════════════════════════
# SAVE
# ════════════════════════════════════════════════════════════════════════════
output_path = r"c:\Users\DeepthiKonchada\Downloads\AutomationNew\testData\MarkAsUsed_TestCases.xlsx"
wb.save(output_path)
print(f"Excel file created: {output_path}")
